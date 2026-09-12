"""Parse and import GBM investment workbook (Excel or Google Sheets rows)."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Optional

import openpyxl

from app.config import GBM_EXCEL_PATH
from app.database import get_db, init_db
from app.history import record_portfolio_snapshot, seed_history_if_empty


def _num(value: Any) -> Optional[float]:
    if value is None or value == "" or value == "#N/A":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").replace("$", "").strip())
    except (TypeError, ValueError):
        return None


def _parse_portfolio_rows(rows: list[list[Any]]) -> dict:
    """Parse Portafolio sheet from row lists (header at row index 2)."""
    holdings = []
    invested = market_value = cash = pnl = return_pct = 0.0

    if rows:
        summary = rows[0] if len(rows) > 0 else []
        for i, cell in enumerate(summary):
            label = str(cell or "").strip().lower()
            if label == "inversión" and i + 1 < len(summary):
                invested = _num(summary[i + 1]) or invested
            elif label == "mercado" and i + 1 < len(summary):
                market_value = _num(summary[i + 1]) or market_value
            elif label == "efectivo" and i + 1 < len(summary):
                cash = _num(summary[i + 1]) or cash
            elif label in ("plus/minus", "plusminus") and i + 1 < len(summary):
                pnl = _num(summary[i + 1]) or pnl
            elif label == "% rendimiento" and i + 1 < len(summary):
                return_pct = _num(summary[i + 1]) or return_pct

    for row in rows[3:]:
        if not row:
            continue
        ticker = row[0] if len(row) > 0 else None
        if not ticker or not str(ticker).startswith("BMV:"):
            continue
        shares = _num(row[5] if len(row) > 5 else None)
        if not shares or shares <= 0:
            continue
        holdings.append({
            "ticker": str(ticker),
            "name": str(row[4] or "") if len(row) > 4 else "",
            "shares": shares,
            "avg_cost": _num(row[6] if len(row) > 6 else None) or 0,
            "market_price": _num(row[7] if len(row) > 7 else None) or 0,
            "market_value": _num(row[9] if len(row) > 9 else None) or 0,
            "pnl": _num(row[10] if len(row) > 10 else None) or 0,
            "weight_pct": (_num(row[12] if len(row) > 12 else None) or 0) * 100,
        })

    if holdings and not market_value:
        market_value = sum(h["market_value"] for h in holdings)
    if holdings and not invested:
        invested = sum(h["avg_cost"] * h["shares"] for h in holdings)
    if not pnl and invested and market_value:
        pnl = market_value - invested
    if invested:
        return_pct = (pnl / invested) * 100
    elif abs(return_pct) < 1:
        return_pct = return_pct * 100

    return {
        "snapshot": {
            "invested": invested,
            "market_value": market_value,
            "cash": cash,
            "pnl": pnl,
            "return_pct": return_pct,
        },
        "holdings": holdings,
    }


def _parse_monitor_rows(rows: list[list[Any]]) -> list[dict]:
    indices = []
    for row in rows[1:]:
        if not row or len(row) < 6:
            continue
        ticker = row[1] if len(row) > 1 else None
        if not ticker or str(ticker) in ("Ticker", "#N/A"):
            continue
        quote = _num(row[5] if len(row) > 5 else None)
        if quote is None:
            continue
        indices.append({
            "ticker": str(ticker),
            "name": str(row[4] or row[3] or "") if len(row) > 4 else "",
            "quote": quote,
            "change_abs": _num(row[6] if len(row) > 6 else None) or 0,
            "change_pct": (_num(row[8] if len(row) > 8 else None) or 0) * 100,
        })
    return indices


def _sheet_to_rows(ws) -> list[list[Any]]:
    return [[ws.cell(r, c).value for c in range(1, ws.max_column + 1)] for r in range(1, ws.max_row + 1)]


def parse_gbm_workbook(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    portfolio = _parse_portfolio_rows(_sheet_to_rows(wb["Portafolio"]))
    indices = _parse_monitor_rows(_sheet_to_rows(wb["Monitor"]))
    return {**portfolio, "indices": indices}


def parse_gbm_sheets(portfolio_rows: list[list[Any]], monitor_rows: list[list[Any]]) -> dict:
    return {
        **_parse_portfolio_rows(portfolio_rows),
        "indices": _parse_monitor_rows(monitor_rows),
    }


def _save_to_db(data: dict, source: str) -> None:
    holdings = data.get("holdings") or []
    if not holdings:
        raise ValueError("El import no contiene posiciones — no se modificó la base de datos")

    now = datetime.now().isoformat(timespec="seconds")
    conn = get_db()

    snap = data["snapshot"]
    conn.execute("DELETE FROM investment_snapshot")
    conn.execute(
        """INSERT INTO investment_snapshot
           (invested, market_value, cash, pnl, return_pct, updated_at, price_source, price_fetched_at)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
        (snap["invested"], snap["market_value"], snap["cash"], snap["pnl"], snap["return_pct"],
         now, source, now),
    )

    imported_tickers: list[str] = []
    for h in holdings:
        imported_tickers.append(h["ticker"])
        existing = conn.execute(
            "SELECT id FROM investment_holdings WHERE ticker = ?", (h["ticker"],)
        ).fetchone()
        if existing:
            conn.execute(
                """UPDATE investment_holdings
                   SET name = ?, shares = ?, avg_cost = ?, market_price = ?, market_value = ?,
                       pnl = ?, weight_pct = ?, updated_at = ?, price_source = ?, price_fetched_at = ?
                   WHERE ticker = ?""",
                (h["name"], h["shares"], h["avg_cost"], h["market_price"],
                 h["market_value"], h["pnl"], h["weight_pct"], now, source, now, h["ticker"]),
            )
        else:
            conn.execute(
                """INSERT INTO investment_holdings
                   (ticker, name, shares, avg_cost, market_price, market_value, pnl, weight_pct,
                    updated_at, price_source, price_fetched_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (h["ticker"], h["name"], h["shares"], h["avg_cost"], h["market_price"],
                 h["market_value"], h["pnl"], h["weight_pct"], now, source, now),
            )
        if h.get("market_price"):
            conn.execute(
                """INSERT INTO price_cache (ticker, price, source, fetched_at)
                   VALUES (?, ?, ?, ?)
                   ON CONFLICT(ticker) DO UPDATE SET
                     price=excluded.price, source=excluded.source, fetched_at=excluded.fetched_at""",
                (h["ticker"], h["market_price"], source, now),
            )

    placeholders = ",".join("?" * len(imported_tickers))
    conn.execute(
        f"""DELETE FROM investment_holdings
            WHERE ticker NOT IN ({placeholders})
              AND COALESCE(price_source, '') != 'manual'""",
        imported_tickers,
    )

    conn.execute("DELETE FROM market_indices")
    for idx in data.get("indices", []):
        conn.execute(
            """INSERT INTO market_indices (ticker, name, quote, change_abs, change_pct, updated_at)
               VALUES (?, ?, ?, ?, ?, ?)""",
            (idx["ticker"], idx["name"], idx["quote"], idx["change_abs"], idx["change_pct"], now),
        )

    conn.execute(
        "INSERT INTO sync_log (source, status, message, synced_at) VALUES (?, ?, ?, ?)",
        (source, "ok", f"{len(data['holdings'])} posiciones importadas", now),
    )

    # Al importar explícitamente, sincronizar GBM al mes calendario actual si existe fila.
    today = datetime.now().date()
    row = conn.execute(
        "SELECT year, month FROM patrimony WHERE year = ? AND month = ?",
        (today.year, today.month),
    ).fetchone()
    if row:
        conn.execute(
            "UPDATE patrimony SET gbm = ? WHERE year = ? AND month = ?",
            (snap["market_value"], today.year, today.month),
        )

    conn.commit()
    conn.close()

    record_portfolio_snapshot(snap["market_value"], snap["invested"], snap["pnl"], snap["return_pct"])
    seed_history_if_empty(snap["market_value"], snap["invested"], snap["pnl"], snap["return_pct"])


def import_from_excel(path: Optional[Path] = None) -> dict:
    init_db()
    excel_path = path or GBM_EXCEL_PATH
    if not excel_path.exists():
        raise FileNotFoundError(f"No se encontró {excel_path}")
    data = parse_gbm_workbook(excel_path)
    _save_to_db(data, "excel")
    return data


def has_investment_data() -> bool:
    conn = get_db()
    count = conn.execute("SELECT COUNT(*) AS n FROM investment_holdings").fetchone()["n"]
    conn.close()
    return count > 0
